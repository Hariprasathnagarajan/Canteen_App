"""
VDart Cafe Automated Test Suite
Covers:
  1. Authentication & Sessions
  2. Role-Based Access Control (RBAC)
  3. Token Allocation & Management
  4. Ordering Transactions & Atomic Balance Deductions
  5. Kitchen Workflow & Automated Token Refunds
  6. Password Reset Security (Self-service + Admin Approval)
  7. Admin Dashboard Analytics & Reporting
"""
from decimal import Decimal
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.utils import timezone
from rest_framework import status
from rest_framework.test import APITestCase, APIClient

from api.models import (
    Office, Kitchen,
    CustomUser, MenuItem, Order, OrderItem,
    ShiftTokenAllocation, TokenDistribution, PasswordResetRequest,
    CanteenSetting, CompanyBill, FinancialTransaction, TokenResetApproval
)
from api.serializers import CustomUserSerializer, CustomUserCreateSerializer


# ---------------------------------------------------------
# Base test setup - shared fixtures for all test classes
# ---------------------------------------------------------
@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class VDartCafeBaseTestCase(APITestCase):
    def setUp(self):
        from django.core.cache import cache
        cache.clear()
        self.client = APIClient()

        self.admin = CustomUser.objects.create_superuser(
            username="admin_test", user_id="ADM001", email="admin@vdartcafe.com",
            password="AdminPassword123!", role="admin",
        )
        self.staff_admin = CustomUser.objects.create_user(
            username="staff_admin_test", user_id="SADM001", email="staffadmin@vdartcafe.com",
            password="StaffAdminPass123!", role="staff_admin",
        )
        self.staff = CustomUser.objects.create_user(
            username="staff_test", user_id="STF001", email="staff@vdartcafe.com",
            password="StaffPassword123!", role="staff",
        )
        self.employee = CustomUser.objects.create_user(
            username="emp_test", user_id="EMP001", email="emp@vdartcafe.com",
            password="EmpPassword123!", role="employee", work_shift="day", monthly_tokens=50,
        )
        self.guest = CustomUser.objects.create_user(
            username="guest_test", user_id="GST001", email="guest@vdartcafe.com",
            password="GuestPassword123!", role="guest", work_shift="day", monthly_tokens=30,
        )
        self.coffee = MenuItem.objects.create(
            name="Test Coffee", description="Fresh hot coffee",
            price=5, is_available=True,
        )
        self.sandwich = MenuItem.objects.create(
            name="Test Sandwich", description="Crispy grilled sandwich",
            price=15, is_available=True,
        )
        self.unavailable = MenuItem.objects.create(
            name="Special Cake", description="Seasonal cake",
            price=20, is_available=False,
        )


# ---------------------------------------------------------
# 1. Authentication & Sessions
# ---------------------------------------------------------
class AuthenticationTests(VDartCafeBaseTestCase):

    def test_login_success(self):
        resp = self.client.post("/api/login/", {"username": "emp_test", "password": "EmpPassword123!"})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["username"], "emp_test")
        self.assertEqual(resp.data["role"], "employee")

    def test_login_wrong_password_rejected(self):
        resp = self.client.post("/api/login/", {"username": "emp_test", "password": "wrong!"})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_profile_requires_authentication(self):
        resp = self.client.get("/api/profile/")
        self.assertIn(resp.status_code, [status.HTTP_401_UNAUTHORIZED, status.HTTP_403_FORBIDDEN])

    def test_profile_returns_correct_user_data(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.get("/api/profile/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["username"], "emp_test")
        self.assertEqual(resp.data["tokens"], 50)

    def test_logout_returns_success(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.post("/api/logout/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)


# ---------------------------------------------------------
# 2. Role-Based Access Control
# ---------------------------------------------------------
class RBACTests(VDartCafeBaseTestCase):

    def test_employee_cannot_access_admin_users_list(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.get("/api/admin/users/")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_admin_can_access_admin_users_list(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/users/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

    def test_staff_can_list_kitchen_orders(self):
        self.client.force_authenticate(user=self.staff)
        resp = self.client.get("/api/staff/orders/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

    def test_employee_cannot_list_kitchen_orders(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.get("/api/staff/orders/")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)


# ---------------------------------------------------------
# 3. Token Allocation & Management
# ---------------------------------------------------------
class TokenAllocationTests(VDartCafeBaseTestCase):

    def test_new_user_starts_with_zero_tokens(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/admin/users/", {
            "username": "night_emp", "user_id": "EMP099", "email": "night@vdartcafe.com",
            "password": "Night123!", "role": "employee", "work_shift": "night",
        })
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        created_user = CustomUser.objects.get(username="night_emp")
        # Tokens are never auto-allocated — the admin allocates explicitly
        self.assertEqual(created_user.monthly_tokens, 0)
        self.assertEqual(created_user.user_id, "EMP099")

    def test_assign_tokens_updates_user_balance_and_creates_audit_trail(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/admin/tokens/assign/", {"shift": "day", "tokens": 65})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 65)
        # Audit trail record exists
        self.assertTrue(TokenDistribution.objects.filter(user=self.employee, tokens_allocated=65).exists())

    def test_assign_tokens_requires_explicit_amount(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/admin/tokens/assign/", {"shift": "day"})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 50)  # unchanged

    def test_assign_all_shifts_requires_explicit_amounts(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/admin/tokens/assign/", {"shift": "all"})
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_admin_token_summary(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/tokens/summary/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn("day_shift", resp.data)
        day_users = resp.data["day_shift"]["users"]
        employee_row = next((u for u in day_users if u["username"] == "emp_test"), None)
        self.assertIsNotNone(employee_row)
        self.assertEqual(employee_row["name"], "emp_test")  # name falls back to username
        self.assertEqual(employee_row["user_id"], "EMP001")
        self.assertIn("email", employee_row)

    def test_assign_tokens_all_shifts(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/admin/tokens/assign/", {
            "shift": "all",
            "tokens": {"day": 50, "mid": 75, "night": 100},
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 50)  # Day shift amount
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.monthly_tokens, 50)  # Guest in day shift


# ---------------------------------------------------------
# 3b. Bulk Import (users + menu) — CSV / JSON / XLSX
# ---------------------------------------------------------
class BulkImportTests(VDartCafeBaseTestCase):

    def _build_xlsx(self, headers, rows):
        from openpyxl import Workbook
        from io import BytesIO

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer.read()

    def test_bulk_import_users_from_pasted_csv(self):
        self.client.force_authenticate(user=self.admin)
        csv_text = (
            "username,user_id,first_name,last_name,email,role,work_shift,monthly_tokens\n"
            "bulk_emp1,BK001,Ravi,Verma,ravi@vdartcafe.com,employee,mid,\n"
            "bulk_emp2,BK002,Sneha,Rao,sneha@vdartcafe.com,employee,night,40\n"
        )
        resp = self.client.post("/api/admin/users/bulk-import/", {"text": csv_text, "format": "csv"})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 2)
        self.assertEqual(resp.data["errors"], [])

        emp1 = CustomUser.objects.get(username="bulk_emp1")
        self.assertEqual(emp1.monthly_tokens, 0)  # no implicit default
        self.assertEqual(emp1.work_shift, "mid")
        emp2 = CustomUser.objects.get(username="bulk_emp2")
        self.assertEqual(emp2.monthly_tokens, 40)
        # Password defaults to user id when not provided
        self.assertTrue(emp1.check_password("BK001"))

    def test_bulk_import_user_without_user_id_defaults_password_to_username(self):
        self.client.force_authenticate(user=self.admin)
        csv_text = (
            "username,user_id,email,role,work_shift\n"
            "no_id_emp,,no_id@vdartcafe.com,employee,day\n"
        )
        resp = self.client.post("/api/admin/users/bulk-import/", {"text": csv_text, "format": "csv"})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)

        emp = CustomUser.objects.get(username="no_id_emp")
        self.assertFalse(emp.user_id)
        # Without a user_id, the password must fall back to the (unique) username
        self.assertTrue(emp.check_password("no_id_emp"))

    def test_bulk_import_users_from_uploaded_csv_file(self):
        self.client.force_authenticate(user=self.admin)
        csv_bytes = (
            "username,user_id,email,role,work_shift\n"
            "file_emp1,FE001,file1@vdartcafe.com,employee,day\n"
        ).encode("utf-8")
        upload = SimpleUploadedFile("users.csv", csv_bytes, content_type="text/csv")
        resp = self.client.post("/api/admin/users/bulk-import/", {"file": upload}, format="multipart")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)
        self.assertTrue(CustomUser.objects.filter(username="file_emp1").exists())

    def test_bulk_import_users_from_xlsx_file(self):
        self.client.force_authenticate(user=self.admin)
        xlsx_bytes = self._build_xlsx(
            ["username", "user_id", "email", "role", "work_shift"],
            [["xlsx_emp1", "XL001", "xlsx1@vdartcafe.com", "employee", "day"]],
        )
        upload = SimpleUploadedFile(
            "users.xlsx", xlsx_bytes,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp = self.client.post("/api/admin/users/bulk-import/", {"file": upload}, format="multipart")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)
        self.assertTrue(CustomUser.objects.filter(username="xlsx_emp1").exists())

    def test_bulk_import_users_xlsx_with_title_row_above_header(self):
        """A decorated sheet (title row above the header) must still import cleanly."""
        self.client.force_authenticate(user=self.admin)
        from openpyxl import Workbook
        from io import BytesIO

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Users"
        sheet.append(["VDart Cafe — User Import"])
        sheet.append([])
        sheet.append(["username", "user_id", "email", "role", "work_shift"])
        sheet.append(["titled_emp1", "TT001", "titled1@vdartcafe.com", "employee", "mid"])
        sheet.append(["titled_emp2", "TT002", "titled2@vdartcafe.com", "guest", "day"])
        instructions = workbook.create_sheet("Instructions")
        instructions.append(["Fill the Users sheet only."])
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        upload = SimpleUploadedFile(
            "VDart_Cafe_Users_Template.xlsx", buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp = self.client.post("/api/admin/users/bulk-import/", {"file": upload}, format="multipart")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 2)
        self.assertEqual(resp.data["errors"], [])
        self.assertTrue(CustomUser.objects.filter(username="titled_emp1").exists())
        self.assertTrue(CustomUser.objects.filter(username="titled_emp2").exists())
        self.assertFalse(CustomUser.objects.filter(username="VDart Cafe — User Import").exists())

    def test_bulk_import_users_without_price_column_is_not_rejected(self):
        """User headers have no price column; header detection must still find them."""
        self.client.force_authenticate(user=self.admin)
        xlsx_bytes = self._build_xlsx(
            ["Username", "First Name", "Last Name", "Role", "Shift"],
            [["hdr_emp1", "Nila", "Das", "employee", "night"]],
        )
        upload = SimpleUploadedFile(
            "users_no_price.xlsx", xlsx_bytes,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp = self.client.post("/api/admin/users/bulk-import/", {"file": upload}, format="multipart")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)
        self.assertTrue(CustomUser.objects.filter(username="hdr_emp1").exists())

    def test_bulk_import_users_updates_existing_and_reports_bad_rows(self):
        self.client.force_authenticate(user=self.admin)
        csv_text = (
            "username,user_id,work_shift\n"
            "emp_test,EMP001,night\n"
            ",BAD001,day\n"
            "bad_shift_emp,BS001,evening\n"
        )
        resp = self.client.post("/api/admin/users/bulk-import/", {"text": csv_text, "format": "csv"})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["updated"], 1)
        self.assertEqual(len(resp.data["errors"]), 2)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.work_shift, "night")

    def test_bulk_import_menu_items_from_pasted_json(self):
        self.client.force_authenticate(user=self.staff)
        payload = {
            "text": '[{"name": "Masala Dosa", "price": 30, "description": "Crispy dosa"},'
                    ' {"name": "Bad Item", "description": "no price"}]',
            "format": "json",
        }
        resp = self.client.post("/api/staff/menu/bulk-import/", payload, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)
        self.assertEqual(len(resp.data["errors"]), 1)
        self.assertEqual(MenuItem.objects.get(name="Masala Dosa").price, 30)

    def test_bulk_import_menu_items_from_xlsx_file(self):
        self.client.force_authenticate(user=self.staff)
        xlsx_bytes = self._build_xlsx(
            ["name", "price", "description", "is_available"],
            [["Rava Kesari", 25, "Sweet kesari", "yes"]],
        )
        upload = SimpleUploadedFile(
            "menu.xlsx", xlsx_bytes,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp = self.client.post("/api/staff/menu/bulk-import/", {"file": upload}, format="multipart")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)
        item = MenuItem.objects.get(name="Rava Kesari")
        self.assertEqual(item.price, 25)
        self.assertTrue(item.is_available)

    def test_bulk_import_requires_admin_for_users(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.post("/api/admin/users/bulk-import/", {"text": "username\nx1"})
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def _build_multi_sheet_xlsx(self):
        from openpyxl import Workbook
        from io import BytesIO

        workbook = Workbook()
        workbook.active.title = "Instructions"
        sheet = workbook.create_sheet("Menu")
        sheet.append(["name", "price", "description"])
        sheet.append(["Sheet2 Dosa", 30, "Crispy"])
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return buffer.read()

    def test_menu_import_reads_data_from_non_active_sheet(self):
        self.client.force_authenticate(user=self.staff)
        upload = SimpleUploadedFile(
            "menu.xlsx", self._build_multi_sheet_xlsx(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp = self.client.post("/api/staff/menu/bulk-import/", {"file": upload}, format="multipart")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)
        self.assertTrue(MenuItem.objects.filter(name="Sheet2 Dosa").exists())

    def test_menu_import_skips_title_rows_above_header(self):
        from openpyxl import Workbook
        from io import BytesIO

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["VDart Cafe Menu - March"])
        sheet.append([])
        sheet.append(["name", "price"])
        sheet.append(["Titled Idli", 12])
        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        self.client.force_authenticate(user=self.staff)
        upload = SimpleUploadedFile(
            "menu.xlsx", buffer.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        resp = self.client.post("/api/staff/menu/bulk-import/", {"file": upload}, format="multipart")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)
        self.assertTrue(MenuItem.objects.filter(name="Titled Idli").exists())

    def test_menu_import_semicolon_delimited_csv(self):
        self.client.force_authenticate(user=self.staff)
        csv_text = "name;price;category\nSemi Tea;8;beverage\n"
        resp = self.client.post("/api/staff/menu/bulk-import/", {"text": csv_text, "format": "csv"}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)
        item = MenuItem.objects.get(name="Semi Tea")
        self.assertEqual(item.price, 8)
        self.assertEqual(item.category, "beverage")

    def test_menu_import_pasted_console_dump(self):
        self.client.force_authenticate(user=self.staff)
        dump = '{disabled: false, name: "Console Tea", price: "15.00", type: "Tea"}'
        resp = self.client.post("/api/staff/menu/bulk-import/", {"text": dump, "format": "csv"}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)
        self.assertTrue(MenuItem.objects.filter(name="Console Tea").exists())

    def test_menu_import_rejects_legacy_xls(self):
        self.client.force_authenticate(user=self.staff)
        upload = SimpleUploadedFile("menu.xls", b"\xd0\xcf\x11\xe0legacy", content_type="application/vnd.ms-excel")
        resp = self.client.post("/api/staff/menu/bulk-import/", {"file": upload}, format="multipart")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("xls", resp.data["error"].lower())


# ---------------------------------------------------------
# 4. Ordering Transactions
# ---------------------------------------------------------
class OrderingTransactionTests(VDartCafeBaseTestCase):

    def test_employee_place_valid_order_deducts_tokens(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.post("/api/employee/order/", {
            "items": [
                {"menu_item_id": self.coffee.id, "quantity": 2},    # 2 * 5 = 10 tokens
                {"menu_item_id": self.sandwich.id, "quantity": 1},  # 1 * 15 = 15 tokens
            ]
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        self.assertEqual(resp.data["total_tokens"], 25)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 25)   # 50 - 25

    def test_order_rejected_on_insufficient_tokens(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.post("/api/employee/order/", {
            "items": [{"menu_item_id": self.sandwich.id, "quantity": 4}]  # 4 * 15 = 60 > 50
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 50)   # unchanged
        self.assertEqual(Order.objects.filter(user=self.employee).count(), 0)

    def test_order_rejected_for_negative_quantity(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.post("/api/employee/order/", {
            "items": [{"menu_item_id": self.coffee.id, "quantity": -1}]
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_order_rejected_for_zero_quantity(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.post("/api/employee/order/", {
            "items": [{"menu_item_id": self.coffee.id, "quantity": 0}]
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_order_rejected_for_unavailable_item(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.post("/api/employee/order/", {
            "items": [{"menu_item_id": self.unavailable.id, "quantity": 1}]
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)

    def test_guest_can_place_order(self):
        self.client.force_authenticate(user=self.guest)
        resp = self.client.post("/api/guest/order/", {
            "items": [{"menu_item_id": self.coffee.id, "quantity": 1}]
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.monthly_tokens, 25)  # 30 - 5


# ---------------------------------------------------------
# 5. Kitchen Workflow & Token Refunds
# ---------------------------------------------------------
class KitchenWorkflowTests(VDartCafeBaseTestCase):

    def setUp(self):
        super().setUp()
        self.order = Order.objects.create(
            user=self.employee, total_tokens=10,
            status="pending",
        )
        OrderItem.objects.create(
            order=self.order, menu_item=self.coffee,
            quantity=2, tokens_per_item=5,
        )
        self.employee.monthly_tokens = 40
        self.employee.save()

    def _update(self, new_status):
        return self.client.patch(
            f"/api/staff/orders/{self.order.id}/update_status/",
            {"status": new_status},
            format="json"
        )

    def test_approve_pending_order(self):
        self.client.force_authenticate(user=self.staff)
        resp = self._update("approved")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "approved")

    def test_mark_approved_order_completed(self):
        self.order.status = "approved"
        self.order.save()
        self.client.force_authenticate(user=self.staff)
        resp = self._update("completed")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.order.refresh_from_db()
        self.assertEqual(self.order.status, "completed")

    def test_decline_pending_order_refunds_tokens_atomically(self):
        self.client.force_authenticate(user=self.staff)
        resp = self._update("declined")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.order.refresh_from_db()
        self.assertTrue(self.order.tokens_refunded)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 50)   # 40 + 10 restored

    def test_double_decline_blocked_no_duplicate_refund(self):
        self.client.force_authenticate(user=self.staff)
        self._update("declined")
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 50)

        resp = self._update("declined")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 50)   # still 50, no extra refund

    def test_cannot_revert_completed_order_to_pending(self):
        self.order.status = "completed"
        self.order.save()
        self.client.force_authenticate(user=self.staff)
        resp = self._update("pending")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)


# ---------------------------------------------------------
# 6. Password Reset Security
# ---------------------------------------------------------
class PasswordResetSecurityTests(VDartCafeBaseTestCase):

    def _reset(self, identifier="emp_test", email="emp@vdartcafe.com", password="FreshPass123!"):
        return self.client.post("/api/password_reset/", {
            "identifier": identifier, "email": email,
            "new_password": password, "confirm_password": password,
        })

    def _submit_request(self, identifier="emp_test", email="emp@vdartcafe.com", reason="Forgot password"):
        return self.client.post("/api/password_reset/request/", {
            "identifier": identifier, "email": email, "reason": reason,
        })

    def _approve(self, user):
        req = PasswordResetRequest.objects.create(user=user, reason="Reset", status="pending")
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post(f"/api/admin/password-resets/{req.id}/approve/")
        self.client.force_authenticate(user=None)
        req.refresh_from_db()
        return req, resp

    # Public self-service reset
    def test_employee_can_self_service_reset(self):
        # Without an approval nothing is applied, and the caller cannot tell.
        resp = self._reset()
        self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED)
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.check_password("EmpPassword123!"))

        # After a Super Admin approves, the user chooses their own password.
        self._approve(self.employee)
        resp = self._reset()
        self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED)
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.check_password("FreshPass123!"))

    def test_admin_cannot_self_service_reset_and_gets_generic_response(self):
        resp = self._reset(identifier="admin_test", email="admin@vdartcafe.com", password="Hacked999!")
        self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED)
        self.assertNotIn("Admin accounts cannot", str(resp.data))
        self.admin.refresh_from_db()
        self.assertTrue(self.admin.check_password("AdminPassword123!"))

    def test_wrong_email_returns_generic_response(self):
        resp = self._reset(email="hacker@wrong.com")
        self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED)
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.check_password("EmpPassword123!"))

    # Admin-approval reset request
    def test_employee_can_submit_reset_request(self):
        resp = self._submit_request()
        self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED)
        self.assertTrue(PasswordResetRequest.objects.filter(user=self.employee, status="pending").exists())

    def test_admin_cannot_submit_reset_request_and_gets_generic_response(self):
        resp = self._submit_request(identifier="admin_test", email="admin@vdartcafe.com", reason="Takeover")
        self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED)
        self.assertNotIn("Admin accounts cannot", str(resp.data))
        self.assertFalse(PasswordResetRequest.objects.filter(user=self.admin).exists())

    def test_admin_approval_grants_permission_without_setting_a_password(self):
        req, resp = self._approve(self.employee)
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(req.status, "approved")
        self.employee.refresh_from_db()
        # The password is NOT reset to the Employee ID - it is not touched at all.
        self.assertTrue(self.employee.check_password("EmpPassword123!"))
        self.assertFalse(self.employee.check_password("EMP001"))
        self.assertTrue(self.employee.password_reset_allowed)
        # No password is disclosed anywhere.
        self.assertNotIn("EmpPassword123!", str(resp.data))
        self.assertNotIn("EMP001", req.admin_notes)

    def test_admin_approval_for_user_without_user_id_grants_permission_only(self):
        no_id_user = CustomUser.objects.create_user(
            username="emp_no_id", password="something",
            email="noid@vdartcafe.com", role="employee", work_shift="day", user_id=None,
        )
        req, resp = self._approve(no_id_user)
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(req.status, "approved")
        no_id_user.refresh_from_db()
        # The username fallback default is gone: the password is untouched.
        self.assertTrue(no_id_user.check_password("something"))
        self.assertFalse(no_id_user.check_password("emp_no_id"))
        self.assertTrue(no_id_user.password_reset_allowed)
        self.assertNotIn("something", str(resp.data))

    def test_admin_rejects_reset_request(self):
        req = PasswordResetRequest.objects.create(user=self.employee, reason="Reset", status="pending")
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post(f"/api/admin/password-resets/{req.id}/reject/", {"notes": "Not now"})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        req.refresh_from_db()
        self.assertEqual(req.status, "rejected")
        self.employee.refresh_from_db()
        self.assertFalse(self.employee.password_reset_allowed)

    def test_cannot_approve_already_approved_request(self):
        req = PasswordResetRequest.objects.create(user=self.employee, reason="Reset", status="approved")
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post(f"/api/admin/password-resets/{req.id}/approve/")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("Cannot approve", resp.data["error"])

    def test_cannot_reject_already_rejected_request(self):
        req = PasswordResetRequest.objects.create(user=self.employee, reason="Reset", status="rejected")
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post(f"/api/admin/password-resets/{req.id}/reject/")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("Cannot reject", resp.data["error"])

    def test_reset_status_endpoint_is_removed(self):
        PasswordResetRequest.objects.create(user=self.employee, reason="Reset", status="pending")
        resp = self.client.get("/api/password_reset/status/?identifier=emp_test")
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)

    # ---- C1 invariants -------------------------------------------------

    def test_public_reset_is_identical_for_unknown_user_wrong_email_and_no_approval(self):
        unknown = self._reset(identifier="does_not_exist", email="ghost@vdartcafe.com")
        wrong_email = self._reset(email="hacker@wrong.com")
        not_approved = self._reset()

        for resp in (unknown, wrong_email, not_approved):
            self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED)
            self.assertEqual(resp.data, unknown.data)

    def test_reset_without_approval_does_not_change_password(self):
        resp = self._reset()
        self.assertEqual(resp.status_code, status.HTTP_202_ACCEPTED)
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.check_password("EmpPassword123!"))
        self.assertFalse(self.employee.password_reset_allowed)

    def test_permission_is_single_use(self):
        self._approve(self.employee)
        first = self._reset(password="FirstPass123!")
        self.assertEqual(first.status_code, status.HTTP_202_ACCEPTED)
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.check_password("FirstPass123!"))
        self.assertFalse(self.employee.password_reset_allowed)

        second = self._reset(password="SecondPass123!")
        self.assertEqual(second.status_code, status.HTTP_202_ACCEPTED)
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.check_password("FirstPass123!"))
        self.assertFalse(self.employee.check_password("SecondPass123!"))

    def test_reset_rejects_weak_password(self):
        self._approve(self.employee)
        resp = self._reset(password="123")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.check_password("EmpPassword123!"))

    def test_staff_admin_cannot_approve_reset_requests(self):
        req = PasswordResetRequest.objects.create(user=self.employee, reason="Reset", status="pending")
        self.client.force_authenticate(user=self.staff_admin)
        list_resp = self.client.get("/api/admin/password-resets/")
        approve_resp = self.client.post(f"/api/admin/password-resets/{req.id}/approve/")
        self.assertEqual(list_resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(approve_resp.status_code, status.HTTP_403_FORBIDDEN)
        req.refresh_from_db()
        self.assertEqual(req.status, "pending")
        self.employee.refresh_from_db()
        self.assertFalse(self.employee.password_reset_allowed)

    def test_admin_reset_password_does_not_return_plaintext(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post(f"/api/admin/users/{self.employee.id}/reset-password/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertNotIn("password\":", str(resp.data).lower().replace("new password", ""))
        self.assertNotIn("EmpPassword123!", str(resp.data))
        self.employee.refresh_from_db()
        # The password itself is untouched - only the permission is granted.
        self.assertTrue(self.employee.check_password("EmpPassword123!"))
        self.assertTrue(self.employee.password_reset_allowed)

    def test_user_serializers_do_not_expose_password_reset_allowed(self):
        self.assertNotIn("password_reset_allowed", CustomUserSerializer().fields)
        self.assertNotIn("password_reset_allowed", CustomUserCreateSerializer().fields)


# ---------------------------------------------------------
# 7. Admin Dashboard Analytics
# ---------------------------------------------------------
class DashboardAnalyticsTests(VDartCafeBaseTestCase):

    def setUp(self):
        super().setUp()
        Order.objects.create(
            user=self.employee, total_tokens=20,
            status="completed",
        )

    def test_dashboard_stats_returns_expected_fields(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/dashboard/stats/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        for field in ["totalUsers", "totalEmployees", "totalMenuItems", "todaysRevenue", "pendingOrders"]:
            self.assertIn(field, resp.data, f"Missing field: {field}")

    def test_dashboard_user_count_includes_all_roles(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/dashboard/stats/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(resp.data["totalUsers"], 4)

    def test_recent_orders_endpoint(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/dashboard/orders/recent/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIsInstance(resp.data, list)

    def test_revenue_data_endpoint(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/dashboard/revenue/?period=week")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIsInstance(resp.data, list)

    def test_non_admin_cannot_access_dashboard(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.get("/api/admin/dashboard/stats/")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)


# ---------------------------------------------------------
# 8. Digital Receipts & Extended Kitchen Workflows
# ---------------------------------------------------------
class DigitalReceiptTests(VDartCafeBaseTestCase):

    def setUp(self):
        super().setUp()
        self.order = Order.objects.create(
            user=self.employee,
            total_tokens=20,
            status="completed",
        )
        self.item1 = OrderItem.objects.create(
            order=self.order,
            menu_item=self.coffee,
            quantity=1,
            tokens_per_item=5,
        )
        self.item2 = OrderItem.objects.create(
            order=self.order,
            menu_item=self.sandwich,
            quantity=1,
            tokens_per_item=15,
        )

    def test_employee_can_fetch_own_order_receipt(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.get(f"/api/orders/{self.order.id}/receipt/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["order_id"], self.order.id)
        self.assertIn("receipt_number", resp.data)
        self.assertTrue(resp.data["receipt_number"].startswith("VDART-RCPT-"))
        self.assertEqual(resp.data["total_tokens"], 20)
        self.assertEqual(len(resp.data["items"]), 2)
        self.assertEqual(resp.data["user"]["username"], "emp_test")
        self.assertEqual(resp.data["user"]["user_id"], "EMP001")
        self.assertIn("verification_code", resp.data)
        self.assertIn("cafe_info", resp.data)
        self.assertEqual(resp.data["cafe_info"]["name"], "VDart Cafe")

    def test_non_owner_forbidden_from_other_user_receipt(self):
        # Guest user attempts to access employee's receipt
        self.client.force_authenticate(user=self.guest)
        resp = self.client.get(f"/api/orders/{self.order.id}/receipt/")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_staff_and_admin_can_view_any_order_receipt(self):
        # Staff inspection for kitchen printing
        self.client.force_authenticate(user=self.staff)
        resp_staff = self.client.get(f"/api/orders/{self.order.id}/receipt/")
        self.assertEqual(resp_staff.status_code, status.HTTP_200_OK)

        # Admin auditing
        self.client.force_authenticate(user=self.admin)
        resp_admin = self.client.get(f"/api/orders/{self.order.id}/receipt/")
        self.assertEqual(resp_admin.status_code, status.HTTP_200_OK)

    def test_receipt_nonexistent_order_returns_404(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.get("/api/orders/999999/receipt/")
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)


class KitchenLiveFeedAndWorkflowTests(VDartCafeBaseTestCase):

    def setUp(self):
        super().setUp()
        # Seed orders with different categories and statuses
        self.pending_order = Order.objects.create(
            user=self.employee, total_tokens=5, status="pending"
        )
        OrderItem.objects.create(
            order=self.pending_order, menu_item=self.coffee, quantity=1, tokens_per_item=5
        )

        self.approved_order = Order.objects.create(
            user=self.guest, total_tokens=15, status="approved"
        )
        OrderItem.objects.create(
            order=self.approved_order, menu_item=self.sandwich, quantity=1, tokens_per_item=15
        )

    def test_staff_live_feed_endpoint(self):
        self.client.force_authenticate(user=self.staff)
        resp = self.client.get("/api/staff/orders/live_feed/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn("counts", resp.data)
        self.assertIn("latest_order_id", resp.data)
        self.assertIn("server_time", resp.data)
        self.assertGreaterEqual(resp.data["counts"]["pending"], 1)
        self.assertGreaterEqual(resp.data["counts"]["approved"], 1)

    def test_staff_orders_category_filtering(self):
        self.client.force_authenticate(user=self.staff)

        # Coffee was auto-categorized as beverage in MenuItem.save()
        resp_bev = self.client.get("/api/staff/orders/?category=beverage")
        self.assertEqual(resp_bev.status_code, status.HTTP_200_OK)
        order_ids = [o["id"] for o in resp_bev.data]
        self.assertIn(self.pending_order.id, order_ids)

    def test_staff_orders_status_filtering(self):
        self.client.force_authenticate(user=self.staff)

        resp_pending = self.client.get("/api/staff/orders/?status=pending")
        self.assertEqual(resp_pending.status_code, status.HTTP_200_OK)
        for o in resp_pending.data:
            self.assertEqual(o["status"], "pending")

        resp_approved = self.client.get("/api/staff/orders/?status=approved")
        self.assertEqual(resp_approved.status_code, status.HTTP_200_OK)
        for o in resp_approved.data:
            self.assertEqual(o["status"], "approved")

    def test_employee_orders_includes_server_time_and_token_balance(self):
        self.client.force_authenticate(user=self.employee)
        resp = self.client.get("/api/employee/orders/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn("today_orders", resp.data)
        self.assertIn("past_orders", resp.data)
        self.assertIn("server_time", resp.data)
        self.assertIn("tokens_balance", resp.data)
        self.assertEqual(resp.data["tokens_balance"], 50)


# ---------------------------------------------------------
# 8. Staff Admin Role & RBAC Boundaries
# ---------------------------------------------------------
class StaffAdminRoleAndRBACBoundaryTests(VDartCafeBaseTestCase):

    def test_staff_admin_can_view_staff_users_only(self):
        self.client.force_authenticate(user=self.staff_admin)
        resp = self.client.get("/api/admin/users/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        for u in resp.data:
            self.assertEqual(u["role"], "staff")

    def test_staff_admin_can_create_staff_user(self):
        self.client.force_authenticate(user=self.staff_admin)
        payload = {
            "username": "new_staff_cook",
            "user_id": "STF099",
            "email": "cook@vdartcafe.com",
            "password": "CookPassword123!",
            "role": "staff",
            "work_shift": "day",
        }
        resp = self.client.post("/api/admin/users/", payload)
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        self.assertEqual(resp.data["role"], "staff")

    def test_staff_admin_cannot_create_admin_or_employee(self):
        self.client.force_authenticate(user=self.staff_admin)
        admin_payload = {
            "username": "rogue_admin",
            "user_id": "ADM999",
            "password": "AdminPassword123!",
            "role": "admin",
        }
        resp = self.client.post("/api/admin/users/", admin_payload)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

        emp_payload = {
            "username": "rogue_emp",
            "user_id": "EMP999",
            "password": "EmpPassword123!",
            "role": "employee",
        }
        resp = self.client.post("/api/admin/users/", emp_payload)
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_staff_admin_cannot_modify_admin_account(self):
        self.client.force_authenticate(user=self.staff_admin)
        resp = self.client.patch(f"/api/admin/users/{self.admin.id}/", {"first_name": "Hacked"})
        self.assertIn(resp.status_code, [status.HTTP_403_FORBIDDEN, status.HTTP_404_NOT_FOUND])

    def test_staff_admin_cannot_reset_password_of_admin_or_employee(self):
        self.client.force_authenticate(user=self.staff_admin)
        resp = self.client.post(f"/api/admin/users/{self.employee.id}/reset-password/", {"new_password": "NewSecretPass123!"})
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)
        self.assertIn("Staff Admins are only permitted", resp.data["error"])

    def test_admin_can_create_staff_admin_user(self):
        self.client.force_authenticate(user=self.admin)
        payload = {
            "username": "new_staff_manager",
            "user_id": "SADM099",
            "email": "manager@vdartcafe.com",
            "role": "staff_admin",
            "work_shift": "day",
        }
        resp = self.client.post("/api/admin/users/", payload)
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        self.assertEqual(resp.data["role"], "staff_admin")
        user = CustomUser.objects.get(username="new_staff_manager")
        self.assertEqual(user.role, "staff_admin")
        self.assertTrue(user.is_staff)

    def test_staff_admin_can_access_revenue_and_dashboard_stats(self):
        self.client.force_authenticate(user=self.staff_admin)
        stats_resp = self.client.get("/api/admin/dashboard/stats/")
        self.assertEqual(stats_resp.status_code, status.HTTP_200_OK)
        self.assertIn("totalUsers", stats_resp.data)

        rev_resp = self.client.get("/api/admin/dashboard/revenue/?period=week")
        self.assertEqual(rev_resp.status_code, status.HTTP_200_OK)
        self.assertIsInstance(rev_resp.data, list)

    def test_staff_admin_can_download_audit_export(self):
        self.client.force_authenticate(user=self.staff_admin)
        csv_resp = self.client.get("/api/reports/audit-export/?format=csv")
        self.assertEqual(csv_resp.status_code, status.HTTP_200_OK)
        self.assertEqual(csv_resp["Content-Type"], "text/csv")



# ---------------------------------------------------------
# 9. Soft Deletion & Historical Order Preservation
# ---------------------------------------------------------
class SoftDeletionAndHistoricalPreservationTests(VDartCafeBaseTestCase):

    def test_soft_deleting_user_preserves_orders_and_receipts(self):
        # Create an order for the employee
        order = Order.objects.create(user=self.employee, total_tokens=5, status="completed")
        OrderItem.objects.create(order=order, menu_item=self.coffee, quantity=1, tokens_per_item=5)

        # Admin soft deletes the employee
        self.client.force_authenticate(user=self.admin)
        resp = self.client.delete(f"/api/admin/users/{self.employee.id}/")
        self.assertEqual(resp.status_code, status.HTTP_204_NO_CONTENT)

        # Verify employee is soft-deleted, not purged from DB
        self.employee.refresh_from_db()
        self.assertTrue(self.employee.is_deleted)
        self.assertFalse(self.employee.is_active)
        self.assertIsNotNone(self.employee.deleted_at)

        # Verify historical order and order item still exist and are linked
        order.refresh_from_db()
        self.assertEqual(order.user_id, self.employee.id)
        self.assertEqual(order.total_tokens, 5)

        # Receipt generator still works cleanly
        self.client.force_authenticate(user=self.admin)
        receipt_resp = self.client.get(f"/api/orders/{order.id}/receipt/")
        self.assertEqual(receipt_resp.status_code, status.HTTP_200_OK)

    def test_soft_deleting_menu_item_preserves_historical_order_items(self):
        order = Order.objects.create(user=self.employee, total_tokens=15, status="completed")
        item = OrderItem.objects.create(order=order, menu_item=self.sandwich, quantity=1, tokens_per_item=15)

        # Staff soft deletes the sandwich
        self.client.force_authenticate(user=self.staff)
        resp = self.client.delete(f"/api/staff/menu/{self.sandwich.id}/")
        self.assertEqual(resp.status_code, status.HTTP_204_NO_CONTENT)

        self.sandwich.refresh_from_db()
        self.assertTrue(self.sandwich.is_deleted)
        self.assertFalse(self.sandwich.is_available)

        # Order item still exists and links to the sandwich
        item.refresh_from_db()
        self.assertEqual(item.menu_item_id, self.sandwich.id)

    def test_token_distribution_protected_against_cascade_deletion(self):
        from django.db.models import ProtectedError
        TokenDistribution.objects.create(user=self.employee, tokens_allocated=50)
        # Directly attempting hard-delete on user raises ProtectedError
        with self.assertRaises(ProtectedError):
            super(CustomUser, self.employee).delete()

    def test_deleted_user_hidden_from_default_list(self):
        self.employee.is_deleted = True
        self.employee.deleted_at = timezone.now()
        self.employee.save()

        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/users/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        data = resp.data["results"] if isinstance(resp.data, dict) else resp.data
        usernames = [u["username"] for u in data]
        self.assertNotIn("emp_test", usernames)

    def test_deleted_user_visible_with_status_filter(self):
        self.employee.is_deleted = True
        self.employee.deleted_at = timezone.now()
        self.employee.save()

        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/users/?status=deleted")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        data = resp.data["results"] if isinstance(resp.data, dict) else resp.data
        usernames = [u["username"] for u in data]
        self.assertIn("emp_test", usernames)

    def test_restore_user_reactivates_and_reappears(self):
        order = Order.objects.create(user=self.employee, total_tokens=5, status="completed")

        self.client.force_authenticate(user=self.admin)
        self.client.delete(f"/api/admin/users/{self.employee.id}/")

        resp = self.client.post(f"/api/admin/users/{self.employee.id}/restore/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertFalse(resp.data["is_deleted"])

        self.employee.refresh_from_db()
        self.assertFalse(self.employee.is_deleted)
        self.assertTrue(self.employee.is_active)
        self.assertIsNone(self.employee.deleted_at)

        # Historical order untouched
        order.refresh_from_db()
        self.assertEqual(order.user_id, self.employee.id)

        # Now visible again in the default list
        list_resp = self.client.get("/api/admin/users/")
        data = list_resp.data["results"] if isinstance(list_resp.data, dict) else list_resp.data
        usernames = [u["username"] for u in data]
        self.assertIn("emp_test", usernames)

    def test_staff_admin_cannot_restore_non_staff(self):
        self.employee.is_deleted = True
        self.employee.deleted_at = timezone.now()
        self.employee.save()

        self.client.force_authenticate(user=self.staff_admin)
        resp = self.client.post(f"/api/admin/users/{self.employee.id}/restore/")
        self.assertEqual(resp.status_code, status.HTTP_403_FORBIDDEN)

    def test_restore_active_user_returns_404(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post(f"/api/admin/users/{self.employee.id}/restore/")
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)



# ---------------------------------------------------------
# 10. Configurable Daily Ordering Limits
# ---------------------------------------------------------
class DailyOrderingLimitsTests(VDartCafeBaseTestCase):

    def setUp(self):
        super().setUp()
        self.setting = CanteenSetting.get_settings()

    def test_daily_order_count_limit_enforced(self):
        self.setting.max_orders_per_user_per_day = 1
        self.setting.save()

        self.client.force_authenticate(user=self.employee)
        payload = {"items": [{"menu_item_id": self.coffee.id, "quantity": 1}]}

        # Order 1: Success
        resp1 = self.client.post("/api/employee/order/", payload, format="json")
        self.assertEqual(resp1.status_code, status.HTTP_201_CREATED)

        # Order 2: Rejected by daily order limit
        resp2 = self.client.post("/api/employee/order/", payload, format="json")
        self.assertEqual(resp2.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("Daily order limit reached", resp2.data["error"])

    def test_daily_item_quantity_limit_enforced(self):
        self.setting.max_orders_per_user_per_day = None
        self.setting.max_qty_per_item_per_day = 2
        self.setting.save()

        self.client.force_authenticate(user=self.employee)
        # Attempt to order 3 coffees when max per item is 2
        payload = {"items": [{"menu_item_id": self.coffee.id, "quantity": 3}]}
        resp = self.client.post("/api/employee/order/", payload, format="json")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("Daily limit exceeded", resp.data["error"])


# ---------------------------------------------------------
# 11. Early Order Cancellation & Token Restoration
# ---------------------------------------------------------
class EarlyOrderCancellationTests(VDartCafeBaseTestCase):

    def test_cancel_pending_order_restores_tokens_immediately(self):
        self.client.force_authenticate(user=self.employee)
        payload = {"items": [{"menu_item_id": self.sandwich.id, "quantity": 1}]}
        place_resp = self.client.post("/api/employee/order/", payload, format="json")
        self.assertEqual(place_resp.status_code, status.HTTP_201_CREATED)
        order_id = place_resp.data["id"]

        # Tokens dropped from 50 to 35
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.current_tokens(), 35)

        # Cancel order
        cancel_resp = self.client.post(f"/api/orders/{order_id}/cancel/", {"reason": "Ordered wrong shift lunch"})
        self.assertEqual(cancel_resp.status_code, status.HTTP_200_OK)
        self.assertEqual(cancel_resp.data["status"], "cancelled")

        # Tokens refunded back to 50
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.current_tokens(), 50)

        order = Order.objects.get(id=order_id)
        self.assertEqual(order.status, "cancelled")
        self.assertTrue(order.tokens_refunded)

    def test_cannot_cancel_approved_or_completed_order(self):
        order = Order.objects.create(user=self.employee, total_tokens=5, status="approved")
        self.client.force_authenticate(user=self.employee)
        resp = self.client.post(f"/api/orders/{order.id}/cancel/")
        self.assertEqual(resp.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("cannot be cancelled", resp.data["error"])

    def test_guest_can_cancel_own_pending_order(self):
        self.client.force_authenticate(user=self.guest)
        payload = {"items": [{"menu_item_id": self.coffee.id, "quantity": 1}]}
        place_resp = self.client.post("/api/guest/order/", payload, format="json")
        self.assertEqual(place_resp.status_code, status.HTTP_201_CREATED)
        order_id = place_resp.data["id"]

        cancel_resp = self.client.post(f"/api/orders/{order_id}/cancel/")
        self.assertEqual(cancel_resp.status_code, status.HTTP_200_OK)
        self.assertEqual(cancel_resp.data["status"], "cancelled")

        order = Order.objects.get(id=order_id)
        self.assertEqual(order.status, "cancelled")
        self.assertTrue(order.tokens_refunded)

    def test_cancellation_idempotent_no_duplicate_refund(self):
        self.client.force_authenticate(user=self.employee)
        payload = {"items": [{"menu_item_id": self.coffee.id, "quantity": 1}]}
        place_resp = self.client.post("/api/employee/order/", payload, format="json")
        order_id = place_resp.data["id"]

        # First cancellation
        resp1 = self.client.post(f"/api/orders/{order_id}/cancel/")
        self.assertEqual(resp1.status_code, status.HTTP_200_OK)

        # Second cancellation retry must be rejected (cannot cancel non-pending)
        resp2 = self.client.post(f"/api/orders/{order_id}/cancel/")
        self.assertEqual(resp2.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("Only pending orders can be cancelled", resp2.data["error"])



# ---------------------------------------------------------
# 12. TXT Menu Importer (DevTools Console Format)
# ---------------------------------------------------------
class TxtMenuImporterTests(VDartCafeBaseTestCase):

    def test_preview_txt_import_from_devtools_dump(self):
        # Create an existing item to verify authentic token preservation
        MenuItem.objects.create(name="Special Filter Tea", price=5, price_inr=18.0, is_available=True)

        self.client.force_authenticate(user=self.staff)
        txt_sample = '''
        {disabled: true, name: "Special Filter Tea", price: "20.00", type: "Tea & Coffee"}
        {disabled: false, name: "Paneer Tikka Roll", price: "120.00", type: "Snacks & Starters"}
        '''
        resp = self.client.post("/api/staff/menu/preview-txt-import/", {"text": txt_sample})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["total_parsed"], 2)
        items = resp.data["items"]
        self.assertEqual(items[0]["name"], "Special Filter Tea")
        self.assertEqual(items[0]["price_inr"], 20.0)
        self.assertEqual(items[0]["tokens"], 5)  # Preserved authentic tokens
        self.assertEqual(items[0]["category"], "beverage")
        self.assertFalse(items[0]["is_available"])

        self.assertEqual(items[1]["name"], "Paneer Tikka Roll")
        self.assertEqual(items[1]["price_inr"], 120.0)
        self.assertIsNone(items[1]["tokens"])  # New item: no synthetic formula
        self.assertEqual(items[1]["category"], "snack")
        self.assertTrue(items[1]["is_available"])

    def test_confirm_txt_import_persists_menu_items(self):
        self.client.force_authenticate(user=self.staff)
        items = [
            {
                "name": "Masala Dosa",
                "price_inr": 80.0,
                "tokens": 8,
                "category": "meal",
                "is_available": True
            }
        ]
        resp = self.client.post("/api/staff/menu/confirm-txt-import/", {"items": items}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp.data["created"], 1)

        dosa = MenuItem.objects.get(name="Masala Dosa")
        self.assertEqual(dosa.price, 8)
        self.assertEqual(float(dosa.price_inr), 80.0)
        self.assertEqual(dosa.category, "meal")

    def test_confirm_txt_import_preserves_existing_token_price(self):
        item = MenuItem.objects.create(name="Samosa", price=15, price_inr=25.0, is_available=True)
        self.client.force_authenticate(user=self.staff)
        items = [{"name": "Samosa", "price_inr": 30.0, "category": "snack", "is_available": True}]
        resp = self.client.post("/api/staff/menu/confirm-txt-import/", {"items": items}, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        item.refresh_from_db()
        self.assertEqual(item.price, 15)  # Preserved authentic token price
        self.assertEqual(float(item.price_inr), 30.0)


# ---------------------------------------------------------
# 13. Dual Financial Accounting & Company Billing
# ---------------------------------------------------------
class DualFinancialAccountingTests(VDartCafeBaseTestCase):

    def test_financial_transaction_crud_and_summary(self):
        self.client.force_authenticate(user=self.staff_admin)
        tx_payload = {
            "transaction_type": "manual_payment",
            "amount_inr": 2500.00,
            "payment_method": "upi",
            "payment_reference": "UPI/VDART/12345",
            "notes": "Direct vendor deposit",
        }
        create_resp = self.client.post("/api/finance/transactions/", tx_payload)
        self.assertEqual(create_resp.status_code, status.HTTP_201_CREATED)
        self.assertEqual(float(create_resp.data["amount_inr"]), 2500.00)

        summary_resp = self.client.get("/api/finance/transactions/summary/")
        self.assertEqual(summary_resp.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(summary_resp.data["total_transactions"], 1)
        self.assertGreaterEqual(summary_resp.data["total_volume_inr"], 2500.00)

    def test_company_bill_generation_and_payment_recording(self):
        Order.objects.create(
            user=self.employee,
            total_tokens=10,
            total_amount_inr=Decimal("150.00"),
            status="completed"
        )
        self.client.force_authenticate(user=self.staff_admin)

        now = timezone.now()
        bill_payload = {
            "year": now.year,
            "month": now.month,
        }
        gen_resp = self.client.post("/api/finance/company-bills/generate/", bill_payload)
        self.assertIn(gen_resp.status_code, [status.HTTP_200_OK, status.HTTP_201_CREATED])
        bill_id = gen_resp.data["id"]
        self.assertEqual(float(gen_resp.data["total_amount_inr"]), 150.00)
        self.assertEqual(gen_resp.data["total_tokens_spent"], 10)

        pay_payload = {
            "amount": gen_resp.data["total_amount_inr"],
            "payment_reference": "CORP-WIRE-9988",
            "payment_method": "bank_transfer"
        }
        pay_resp = self.client.post(f"/api/finance/company-bills/{bill_id}/record-payment/", pay_payload)
        self.assertEqual(pay_resp.status_code, status.HTTP_200_OK)
        self.assertEqual(pay_resp.data["bill"]["status"], "paid")

        ft = FinancialTransaction.objects.filter(company_bill_id=bill_id).first()
        self.assertIsNotNone(ft)
        self.assertEqual(ft.payment_reference, "CORP-WIRE-9988")
        self.assertEqual(ft.amount_tokens, 150)
        self.assertNotIn("₹", str(ft))
        self.assertIn("150 tokens", str(ft))

    def test_transaction_is_token_denominated_without_inr(self):
        self.client.force_authenticate(user=self.staff_admin)
        tx_payload = {
            "transaction_type": "manual_payment",
            "amount_tokens": 40,
            "payment_method": "token_transfer",
            "notes": "Shift token quota topup",
        }
        create_resp = self.client.post("/api/finance/transactions/", tx_payload)
        self.assertEqual(create_resp.status_code, status.HTTP_201_CREATED)
        self.assertEqual(create_resp.data["amount_tokens"], 40)

        tx = FinancialTransaction.objects.get(id=create_resp.data["id"])
        self.assertEqual(tx.amount_tokens, 40)
        self.assertNotIn("₹", str(tx))
        self.assertIn("40 tokens", str(tx))
        self.assertNotIn("INR", str(tx))


# ---------------------------------------------------------
# 14. Monthly Audit Export & Formula Injection Defense
# ---------------------------------------------------------
class MonthlyAuditExportTests(VDartCafeBaseTestCase):

    def test_csv_audit_export_sanitizes_formula_injection(self):
        malicious_dish = MenuItem.objects.create(
            name="=CMD|' /C calc'!A0", price=5, is_available=True
        )
        order = Order.objects.create(user=self.employee, total_tokens=5, status="completed")
        OrderItem.objects.create(order=order, menu_item=malicious_dish, quantity=1, tokens_per_item=5)

        self.client.force_authenticate(user=self.staff_admin)
        resp = self.client.get("/api/reports/audit-export/?format=csv")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp["Content-Type"], "text/csv")
        csv_content = resp.content.decode("utf-8")
        self.assertIn("'=CMD", csv_content)
        self.assertNotIn("Total INR", csv_content)
        self.assertNotIn("INR", csv_content)

    def test_xlsx_audit_export_returns_valid_excel_stream(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/reports/audit-export/?format=xlsx")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        self.assertIn("VDart_Audit_", resp["Content-Disposition"])

    def test_pdf_audit_export_returns_valid_reportlab_stream(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/reports/audit-export/?format=pdf")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertEqual(resp["Content-Type"], "application/pdf")
        self.assertTrue(resp.content.startswith(b"%PDF-"))
        self.assertIn("VDart_Audit_", resp["Content-Disposition"])


# ---------------------------------------------------------
# 15. Veg & Non-Veg Food Management & Filtering
# ---------------------------------------------------------
class VegNonVegMenuTests(VDartCafeBaseTestCase):

    def test_veg_nonveg_menu_item_creation_and_autodetection(self):
        # Auto-detection from non-veg keyword
        chicken_dish = MenuItem.objects.create(name="Chicken Biryani", price=25, price_inr=150.0)
        self.assertFalse(chicken_dish.is_veg)
        self.assertEqual(chicken_dish.food_type, "non_veg")

        # Explicit veg dish
        paneer_dish = MenuItem.objects.create(name="Paneer Butter Masala", price=20, price_inr=120.0, is_veg=True, food_type="veg")
        self.assertTrue(paneer_dish.is_veg)
        self.assertEqual(paneer_dish.food_type, "veg")

    def test_staff_can_create_and_update_veg_and_nonveg_dishes(self):
        self.client.force_authenticate(user=self.staff)
        payload = {
            "name": "Egg Fried Rice",
            "price": 18,
            "price_inr": 110.0,
            "category": "meal",
            "is_veg": False,
            "food_type": "non_veg"
        }
        resp = self.client.post("/api/staff/menu/", payload)
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        self.assertFalse(resp.data["is_veg"])
        self.assertEqual(resp.data["food_type"], "non_veg")

    def test_menu_filtering_by_veg_and_nonveg(self):
        MenuItem.objects.create(name="Gobi Manchurian", price=10, is_veg=True, food_type="veg")
        MenuItem.objects.create(name="Fish Curry", price=25, is_veg=False, food_type="non_veg")

        self.client.force_authenticate(user=self.employee)
        # Filter veg
        veg_resp = self.client.get("/api/employee/menu/?food_type=veg")
        self.assertEqual(veg_resp.status_code, status.HTTP_200_OK)
        for item in veg_resp.data:
            self.assertTrue(item["is_veg"])

        # Filter non_veg
        nonveg_resp = self.client.get("/api/employee/menu/?food_type=non_veg")
        self.assertEqual(nonveg_resp.status_code, status.HTTP_200_OK)
        for item in nonveg_resp.data:
            self.assertFalse(item["is_veg"])


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class TokenTrackingAndAutoDistributionTests(APITestCase):
    """
    Tests for:
    - Automatic monthly token distribution service and lazy evaluation
    - Admin tracking of distributed, spent, and unused tokens per diner
    - Individual user transaction tracking audit endpoint
    """

    def setUp(self):
        self.admin = CustomUser.objects.create_superuser(
            username="admin_tracker",
            password="adminpassword123",
            email="tracker_admin@example.com",
            role="admin",
            user_id="ADM-TRK-01"
        )
        self.staff_admin = CustomUser.objects.create_user(
            username="staff_admin_tracker",
            password="staffadminpassword123",
            email="tracker_staffadmin@example.com",
            role="staff_admin",
            user_id="STAFF-ADM-01"
        )
        self.employee = CustomUser.objects.create_user(
            username="emp_tracker",
            password="emppassword123",
            email="emp_tracker@example.com",
            role="employee",
            work_shift="day",
            monthly_tokens=50,
            user_id="EMP-TRK-01"
        )
        self.menu_item = MenuItem.objects.create(
            name="Veg Thali Meal",
            price=15,
            food_type="veg",
            is_veg=True,
            is_available=True
        )

    def test_distribute_approved_monthly_tokens_service(self):
        from .token_service import distribute_approved_monthly_tokens, get_current_cycle_start
        from .models import TokenDistribution, FinancialTransaction
        from datetime import date

        cycle_start = get_current_cycle_start()
        # Configure an explicit shift allocation WITHOUT triggering its per-user
        # distribution side-effect (no implicit default exists).
        ShiftTokenAllocation.objects.bulk_create([
            ShiftTokenAllocation(shift='day', allocation_month=cycle_start, tokens_per_user=100)
        ])
        # Clear out any prior distributions for test user
        TokenDistribution.objects.filter(user=self.employee, allocation_month=cycle_start).delete()
        self.employee.last_token_reset = date(2025, 1, 1)
        self.employee.save()

        # Run the approval-gated distribution
        res = distribute_approved_monthly_tokens(admin_user=self.admin)
        self.assertEqual(res['status'], 'success')
        self.assertGreaterEqual(res['total_tokens_distributed'], 100)

        # Verify DB records - the canonical cycle key is used everywhere
        dist = TokenDistribution.objects.filter(user=self.employee, allocation_month=cycle_start).first()
        self.assertIsNotNone(dist)
        self.assertEqual(dist.tokens_allocated, 100)

        alloc = ShiftTokenAllocation.objects.filter(shift='day', allocation_month=cycle_start).first()
        self.assertIsNotNone(alloc)
        self.assertEqual(alloc.tokens_per_user, 100)

        tx = FinancialTransaction.objects.filter(user=self.employee, payment_method='token_transfer').first()
        self.assertIsNotNone(tx)
        self.assertEqual(tx.amount_tokens, 100)

        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 100)
        self.assertEqual(self.employee.last_token_reset, cycle_start)

        # A second run does not double-allocate
        res2 = distribute_approved_monthly_tokens(admin_user=self.admin)
        self.assertEqual(res2['total_users_allocated'], 0)

    def test_ensure_user_monthly_tokens_lazy_evaluation(self):
        from .token_service import ensure_user_monthly_tokens, get_current_cycle_start
        from datetime import date

        cycle_start = get_current_cycle_start()
        # Configure an explicit shift allocation WITHOUT triggering its per-user
        # distribution side-effect, so we can exercise the lazy top-up path.
        ShiftTokenAllocation.objects.bulk_create([
            ShiftTokenAllocation(shift='day', allocation_month=cycle_start, tokens_per_user=100)
        ])

        # Set user's last reset to a past cycle
        self.employee.last_token_reset = date(2025, 1, 1)
        self.employee.monthly_tokens = 5
        self.employee.save()

        # Without an approved reset nothing is credited (the approval gate holds)
        self.assertFalse(ensure_user_monthly_tokens(self.employee))
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 5)

        # With an approved reset for this cycle the user is credited
        TokenResetApproval.objects.create(
            cycle_month=cycle_start,
            status='approved',
            distributed_tokens=0,
            affected_users_count=0,
            approved_by=self.admin,
        )
        updated = ensure_user_monthly_tokens(self.employee)
        self.assertTrue(updated)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 100)
        self.assertEqual(self.employee.last_token_reset, cycle_start)

    def test_get_token_summary_includes_distributed_spent_and_unused(self):
        from .token_service import get_current_cycle_start
        from .models import TokenDistribution

        cycle_start = get_current_cycle_start()
        TokenDistribution.objects.create(
            user=self.employee,
            allocation_month=cycle_start,
            tokens_allocated=50
        )
        self.employee.last_token_reset = cycle_start

        # Place an order for employee
        order = Order.objects.create(
            user=self.employee,
            status='completed',
            total_tokens=15,
            total_amount_inr=15.0
        )
        OrderItem.objects.create(
            order=order,
            menu_item=self.menu_item,
            quantity=1,
            tokens_per_item=15,
            price_inr_per_item=15.0
        )
        self.employee.monthly_tokens = 35
        self.employee.save()

        self.client.force_authenticate(user=self.admin)
        resp = self.client.get('/api/admin/tokens/summary/')
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

        # Check shift users
        day_users = resp.data['day_shift']['users']
        emp_data = next((u for u in day_users if u['username'] == 'emp_tracker'), None)
        self.assertIsNotNone(emp_data)
        self.assertEqual(emp_data['tokens_distributed'], 50)
        self.assertEqual(emp_data['tokens_spent'], 15)
        self.assertEqual(emp_data['tokens_unused'], 35)

        # Check overall summary
        overall = resp.data['overall']
        self.assertIn('total_distributed', overall)
        self.assertIn('total_spent', overall)
        self.assertIn('total_unused', overall)
        self.assertGreaterEqual(overall['total_spent'], 15)

    def test_auto_distribute_endpoint_is_removed(self):
        self.client.force_authenticate(user=self.staff_admin)
        resp = self.client.post('/api/admin/tokens/auto-distribute/', {'force': True})
        self.assertEqual(resp.status_code, status.HTTP_404_NOT_FOUND)

    def test_user_token_transactions_audit_endpoint(self):
        from .token_service import get_current_cycle_start
        from .models import TokenDistribution

        TokenDistribution.objects.update_or_create(
            user=self.employee,
            allocation_month=get_current_cycle_start(),
            defaults={'tokens_allocated': 50},
        )
        order = Order.objects.create(
            user=self.employee,
            status='completed',
            total_tokens=30,
            total_amount_inr=30.0
        )
        OrderItem.objects.create(
            order=order,
            menu_item=self.menu_item,
            quantity=2,
            tokens_per_item=15,
            price_inr_per_item=15.0
        )
        self.employee.monthly_tokens = 20
        self.employee.save()

        self.client.force_authenticate(user=self.admin)
        resp = self.client.get(f'/api/admin/tokens/user-transactions/{self.employee.id}/')
        self.assertEqual(resp.status_code, status.HTTP_200_OK)

        # Verify user details
        self.assertEqual(resp.data['user']['username'], 'emp_tracker')
        self.assertEqual(resp.data['user']['user_id'], 'EMP-TRK-01')

        # Verify metrics
        metrics = resp.data['metrics']
        self.assertEqual(metrics['tokens_distributed'], 50)
        self.assertEqual(metrics['tokens_spent'], 30)
        self.assertEqual(metrics['tokens_unused'], 20)

        # Verify orders breakdown
        orders = resp.data['orders']
        self.assertEqual(len(orders), 1)
        self.assertEqual(orders[0]['total_tokens'], 30)
        self.assertEqual(len(orders[0]['items']), 1)
        self.assertEqual(orders[0]['items'][0]['name'], 'Veg Thali Meal')
        self.assertEqual(orders[0]['items'][0]['quantity'], 2)

    # ---- Phase 0 regression tests ----------------------------------------

    def test_cycle_key_is_consistent_after_approval(self):
        """N1/N2: every token record is keyed by the cycle start, never the 1st."""
        from .token_service import distribute_approved_monthly_tokens, get_current_cycle_start

        cycle_start = get_current_cycle_start()
        ShiftTokenAllocation.objects.create(
            shift='day', allocation_month=cycle_start, tokens_per_user=120
        )
        distribute_approved_monthly_tokens(admin_user=self.admin)

        dist_keys = set(TokenDistribution.objects.values_list('allocation_month', flat=True))
        alloc_keys = set(ShiftTokenAllocation.objects.values_list('allocation_month', flat=True))
        self.assertEqual(dist_keys, {cycle_start})
        self.assertEqual(alloc_keys, {cycle_start})

        self.assertEqual(
            TokenDistribution.objects.get(user=self.employee).allocation_month, cycle_start
        )
        self.assertEqual(
            ShiftTokenAllocation.objects.get(shift='day').allocation_month, cycle_start
        )

    def test_approve_reset_does_not_raise_integrity_error(self):
        """N1: approving twice for the same cycle must not raise."""
        from .token_service import distribute_approved_monthly_tokens, get_current_cycle_start

        cycle_start = get_current_cycle_start()
        ShiftTokenAllocation.objects.create(
            shift='day', allocation_month=cycle_start, tokens_per_user=120
        )
        distribute_approved_monthly_tokens(admin_user=self.admin)
        # A second approve for the same cycle updates the single approval row.
        distribute_approved_monthly_tokens(admin_user=self.admin)

        approvals = TokenResetApproval.objects.filter(cycle_month=cycle_start)
        self.assertEqual(approvals.count(), 1)
        self.assertEqual(
            TokenDistribution.objects.filter(user=self.employee, allocation_month=cycle_start).count(), 1
        )

    def test_new_user_is_credited_after_approval(self):
        """N3: a diner who joins after the approval is credited on their next action."""
        from .token_service import ensure_user_monthly_tokens, get_current_cycle_start

        cycle_start = get_current_cycle_start()
        ShiftTokenAllocation.objects.create(
            shift='day', allocation_month=cycle_start, tokens_per_user=120
        )
        TokenResetApproval.objects.create(
            cycle_month=cycle_start,
            status='approved',
            distributed_tokens=0,
            affected_users_count=0,
            approved_by=self.admin,
        )

        late_joiner = CustomUser.objects.create_user(
            username="late_joiner", password="latepassword123",
            email="late@example.com", role="employee", work_shift="day",
            user_id="EMP-LATE-01",
        )
        self.assertTrue(ensure_user_monthly_tokens(late_joiner))
        late_joiner.refresh_from_db()
        self.assertEqual(late_joiner.monthly_tokens, 120)
        self.assertEqual(late_joiner.last_token_reset, cycle_start)

    def test_past_cycle_metrics_match_allocated_total(self):
        """N6: the past-cycle window matches the previous cycle's allocations exactly."""
        from .token_service import get_past_cycle_token_metrics, get_billing_cycle_dates
        from datetime import timedelta

        cycle_info = get_billing_cycle_dates()
        prev_start = cycle_info['prev_cycle_start']

        # A past reset must have happened for previous-cycle metrics to be meaningful.
        self.employee.last_token_reset = cycle_info['cycle_start']
        self.employee.save(update_fields=['last_token_reset'])

        TokenDistribution.objects.create(
            user=self.employee, allocation_month=prev_start, tokens_allocated=75
        )
        # A row outside the previous cycle must not be counted.
        TokenDistribution.objects.create(
            user=self.employee,
            allocation_month=prev_start - timedelta(days=1),
            tokens_allocated=500,
        )

        metrics = get_past_cycle_token_metrics()
        self.assertEqual(metrics['past_total_allocated'], 75)


# ---------------------------------------------------------
# 16. Shift-Wise Daily Token Limits & Super Admin Authorization
# ---------------------------------------------------------
class ShiftDailyTokenLimitsTests(VDartCafeBaseTestCase):

    def setUp(self):
        super().setUp()
        self.setting = CanteenSetting.get_settings()
        self.setting.day_shift_daily_limit = 300
        self.setting.mid_shift_daily_limit = 250
        self.setting.night_shift_daily_limit = 300
        self.setting.save()

        # Create dish
        self.item_100tk = MenuItem.objects.create(
            name="Deluxe Feast",
            price=100,
            price_inr=100.0,
            is_available=True
        )

        # Create test diners for each shift
        self.day_user = CustomUser.objects.create_user(
            username="day_shift_worker",
            password="password123",
            role="employee",
            work_shift="day",
            monthly_tokens=1000
        )
        self.mid_user = CustomUser.objects.create_user(
            username="mid_shift_worker",
            password="password123",
            role="employee",
            work_shift="mid",
            monthly_tokens=1000
        )
        self.night_user = CustomUser.objects.create_user(
            username="night_shift_worker",
            password="password123",
            role="employee",
            work_shift="night",
            monthly_tokens=1000
        )

    def test_day_shift_daily_limit_enforcement(self):
        self.client.force_authenticate(user=self.day_user)

        # 300 tk order should succeed (3 x 100)
        resp1 = self.client.post('/api/employee/order/', {
            'items': [{'menu_item_id': self.item_100tk.id, 'quantity': 3}]
        }, format='json')
        self.assertEqual(resp1.status_code, status.HTTP_201_CREATED)

        # Another 100 tk order in the same shift exceeds 300 limit (300 + 100 = 400 > 300)
        resp2 = self.client.post('/api/employee/order/', {
            'items': [{'menu_item_id': self.item_100tk.id, 'quantity': 1}]
        }, format='json')
        self.assertEqual(resp2.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('Daily limit reached for Day Shift', resp2.data['error'])

    def test_mid_shift_daily_limit_enforcement(self):
        self.client.force_authenticate(user=self.mid_user)

        # 200 tk order succeeds (2 x 100 <= 250)
        resp1 = self.client.post('/api/employee/order/', {
            'items': [{'menu_item_id': self.item_100tk.id, 'quantity': 2}]
        }, format='json')
        self.assertEqual(resp1.status_code, status.HTTP_201_CREATED)

        # Trying to order another 100 tk (200 + 100 = 300 > 250) should fail
        resp2 = self.client.post('/api/employee/order/', {
            'items': [{'menu_item_id': self.item_100tk.id, 'quantity': 1}]
        }, format='json')
        self.assertEqual(resp2.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('Daily limit reached for Mid Shift', resp2.data['error'])

    def test_night_shift_daily_limit_enforcement(self):
        self.client.force_authenticate(user=self.night_user)

        # 300 tk order succeeds (3 x 100 <= 300)
        resp1 = self.client.post('/api/employee/order/', {
            'items': [{'menu_item_id': self.item_100tk.id, 'quantity': 3}]
        }, format='json')
        self.assertEqual(resp1.status_code, status.HTTP_201_CREATED)

        # Another 100 tk order exceeds night shift limit (300 + 100 = 400 > 300)
        resp2 = self.client.post('/api/employee/order/', {
            'items': [{'menu_item_id': self.item_100tk.id, 'quantity': 1}]
        }, format='json')
        self.assertEqual(resp2.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('Daily limit reached for Night Shift', resp2.data['error'])

    def test_super_admin_exclusive_settings_modification(self):
        # Staff Admin attempting to modify canteen settings receives 403 Forbidden
        self.client.force_authenticate(user=self.staff_admin)
        resp_forbidden = self.client.patch('/api/admin/settings/', {
            'day_shift_daily_limit': 350
        }, format='json')
        self.assertEqual(resp_forbidden.status_code, status.HTTP_403_FORBIDDEN)
        self.assertIn('Only Super Admin', resp_forbidden.data['error'])

        # Super Admin successfully modifies canteen settings
        self.client.force_authenticate(user=self.admin)
        resp_success = self.client.patch('/api/admin/settings/', {
            'day_shift_daily_limit': 350,
            'mid_shift_daily_limit': 280,
            'night_shift_daily_limit': 320
        }, format='json')
        self.assertEqual(resp_success.status_code, status.HTTP_200_OK)
        self.assertEqual(resp_success.data['day_shift_daily_limit'], 350)
        self.assertEqual(resp_success.data['mid_shift_daily_limit'], 280)
        self.assertEqual(resp_success.data['night_shift_daily_limit'], 320)


# ---------------------------------------------------------
# 10. Monthly Token Allocation Workflow & History Tests
# ---------------------------------------------------------
class MonthlyTokenAllocationWorkflowTests(VDartCafeBaseTestCase):

    def setUp(self):
        super().setUp()
        from zoneinfo import ZoneInfo
        self.ZoneInfo = ZoneInfo
        self.setting = CanteenSetting.get_settings()
        self.setting.monthly_token_allocation_mode = 'automatic'
        self.setting.token_allocation_day = 20
        self.setting.day_shift_policy_quota = 100
        self.setting.mid_shift_policy_quota = 100
        self.setting.night_shift_policy_quota = 100
        self.setting.save()

    def test_01_automatic_mode_calculates_proposal_for_correct_cycle(self):
        """1. Automatic mode ON calculates proposal for correct cycle."""
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/tokens/summary/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn("reset_approval_info", resp.data)
        approval_info = resp.data["reset_approval_info"]
        self.assertEqual(approval_info["allocation_mode"], "automatic")
        self.assertIn("proposal", approval_info)
        proposal = approval_info["proposal"]
        self.assertGreaterEqual(proposal["total_eligible_diners"], 1)
        self.assertIn("shift_breakdown", proposal)
        day_shift = proposal["shift_breakdown"]["day"]
        self.assertEqual(day_shift["policy_quota"], 100)

    def test_02_required_approval_enforced_before_wallet_credits(self):
        """2. Required approval enforced before wallet credits."""
        from api.token_service import get_billing_cycle_dates
        initial_balance = self.employee.monthly_tokens
        cycle_month = get_billing_cycle_dates()["cycle_start"]
        TokenResetApproval.objects.filter(cycle_month=cycle_month).delete()
        TokenResetApproval.objects.create(
            cycle_month=cycle_month,
            status="pending",
            affected_users_count=1,
            distributed_tokens=0,
            allocation_mode="automatic"
        )
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, initial_balance)

        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/admin/tokens/approve-reset/", {
            "quotas": {"day": 80, "mid": 80, "night": 80},
            "notes": "Approved for cycle"
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 80)

    def test_03_automatic_mode_off_prevents_scheduled_and_allows_manual(self):
        """3. Automatic mode OFF prevents scheduled allocation and allows manual allocation."""
        self.setting.monthly_token_allocation_mode = "manual"
        self.setting.save()

        from django.core.management import call_command
        from io import StringIO
        out = StringIO()
        call_command("distribute_monthly_tokens", stdout=out)
        self.assertIn("Automatic monthly token allocation is currently OFF", out.getvalue())

        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/admin/tokens/assign/", {"shift": "day", "tokens": 90})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 90)

    def test_04_toggling_mode_does_not_immediately_credit_tokens(self):
        """4. Toggling mode does not immediately credit tokens."""
        initial_tokens = self.employee.monthly_tokens
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/admin/tokens/mode/", {"mode": "manual"})
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, initial_tokens)

        resp2 = self.client.post("/api/admin/tokens/mode/", {"mode": "automatic"})
        self.assertEqual(resp2.status_code, status.HTTP_200_OK)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, initial_tokens)

    def test_05_previous_cycle_amounts_never_appear_as_new_cycle_inputs(self):
        """5. Previous-cycle amounts never appear as new-cycle actual input values."""
        from api.token_service import get_billing_cycle_dates, get_shift_monthly_quota
        ShiftTokenAllocation.objects.create(
            allocation_month="2026-08-20",
            shift="day",
            tokens_per_user=175
        )
        self.setting.day_shift_policy_quota = 100
        self.setting.save()

        cycle_month = get_billing_cycle_dates()["cycle_start"]
        quota = get_shift_monthly_quota("day", cycle_month)
        self.assertEqual(quota, 100)

    def test_06_previous_cycle_remains_in_history_view_after_new_cycle_begins(self):
        """6. Previous cycle remains in history view after new cycle begins."""
        TokenResetApproval.objects.create(
            cycle_month="2026-08-20",
            cycle_end="2026-09-19",
            status="approved",
            affected_users_count=5,
            distributed_tokens=500,
            allocation_mode="automatic",
            approved_by=self.admin
        )
        TokenResetApproval.objects.create(
            cycle_month="2026-09-20",
            cycle_end="2026-10-19",
            status="approved",
            affected_users_count=5,
            distributed_tokens=500,
            allocation_mode="automatic",
            approved_by=self.admin
        )
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/tokens/history/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        cycle_starts = [str(c["cycle_start"]) for c in resp.data["cycles"]]
        self.assertIn("2026-08-20", cycle_starts)
        self.assertIn("2026-09-20", cycle_starts)

    def test_07_proposed_values_and_actual_credited_values_remain_distinct(self):
        """7. Proposed values and actual credited values remain distinct."""
        from api.token_service import get_billing_cycle_dates
        self.client.force_authenticate(user=self.admin)
        resp = self.client.post("/api/admin/tokens/approve-reset/", {
            "quotas": {"day": 120, "mid": 120, "night": 120},
            "notes": "Custom increased budget"
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        cycle_month = get_billing_cycle_dates()["cycle_start"]
        approval = TokenResetApproval.objects.get(cycle_month=cycle_month)
        self.assertIsNotNone(approval.draft_quotas)
        self.assertEqual(approval.draft_quotas.get("day"), 120)
        self.assertIsNotNone(approval.policy_snapshot)
        self.assertEqual(approval.policy_snapshot.get("day"), 100)

    def test_08_late_joining_employee_handled_by_saved_policy(self):
        """8. Late joining employee / shift change handled by saved policy."""
        self.setting.mid_shift_policy_quota = 110
        self.setting.save()
        new_emp = CustomUser.objects.create_user(
            username="late_joiner", user_id="EMP888", email="late@vdartcafe.com",
            password="LatePassword123!", role="employee", work_shift="mid"
        )
        self.assertEqual(self.setting.get_shift_policy_quota("mid"), 110)
        self.assertEqual(new_emp.monthly_tokens, 0)

    def test_09_scheduler_replay_idempotency_does_not_duplicate_credits(self):
        """9. Scheduler / API replay does not duplicate credits (idempotency)."""
        from api.token_service import distribute_approved_monthly_tokens
        self.client.force_authenticate(user=self.admin)
        self.client.post("/api/admin/tokens/approve-reset/", {
            "quotas": {"day": 70, "mid": 70, "night": 70}
        }, format="json")
        self.employee.refresh_from_db()
        balance_after_first = self.employee.monthly_tokens

        distribute_approved_monthly_tokens(admin_user=self.admin)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, balance_after_first)

    def test_10_safe_retry_of_allocation_transactions(self):
        """10. Safe retry of allocation transactions."""
        from api.token_service import distribute_approved_monthly_tokens
        result = distribute_approved_monthly_tokens(admin_user=self.admin)
        self.assertIn("total_tokens_distributed", result)
        result2 = distribute_approved_monthly_tokens(admin_user=self.admin)
        self.assertIn("total_tokens_distributed", result2)

    def test_11_unauthorized_users_blocked_from_changing_mode_or_approving(self):
        """11. Unauthorized users blocked from changing mode or approving."""
        self.client.force_authenticate(user=self.employee)
        resp1 = self.client.post("/api/admin/tokens/mode/", {"mode": "manual"})
        self.assertEqual(resp1.status_code, status.HTTP_403_FORBIDDEN)

        resp2 = self.client.post("/api/admin/tokens/approve-reset/", {})
        self.assertEqual(resp2.status_code, status.HTTP_403_FORBIDDEN)

        resp3 = self.client.post("/api/admin/tokens/reject-reset/", {})
        self.assertEqual(resp3.status_code, status.HTTP_403_FORBIDDEN)

    def test_12_cycle_dates_across_month_and_year_boundaries(self):
        """12. 20th-to-19th cycle across month and year boundaries."""
        from api.token_service import get_billing_cycle_dates
        pre_jan = timezone.datetime(2027, 1, 15, 10, 0, tzinfo=self.ZoneInfo("Asia/Kolkata"))
        c1 = get_billing_cycle_dates(reference_date=pre_jan)
        self.assertEqual(str(c1["cycle_start"]), "2026-12-20")
        self.assertEqual(str(c1["cycle_end"]), "2027-01-19")

        post_dec = timezone.datetime(2026, 12, 22, 10, 0, tzinfo=self.ZoneInfo("Asia/Kolkata"))
        c2 = get_billing_cycle_dates(reference_date=post_dec)
        self.assertEqual(str(c2["cycle_start"]), "2026-12-20")
        self.assertEqual(str(c2["cycle_end"]), "2027-01-19")

    def test_13_historical_exports_match_persisted_records(self):
        """13. Historical exports match persisted records."""
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/tokens/history/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        self.assertIn("cycles", resp.data)
        self.assertIn("audit_logs", resp.data)

    def test_14_existing_ordering_and_spending_features_unaffected(self):
        """14. Existing ordering and spending features unaffected."""
        self.employee.monthly_tokens = 50
        self.employee.save()
        self.client.force_authenticate(user=self.employee)
        resp = self.client.post("/api/employee/order/", {
            "items": [{"menu_item_id": self.coffee.id, "quantity": 2}]
        }, format="json")
        self.assertEqual(resp.status_code, status.HTTP_201_CREATED)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, 40)


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])
class AdminDashboardCycleAndKPITests(APITestCase):
    def setUp(self):
        self.admin = CustomUser.objects.create_user(
            username="admin_dash_audit",
            email="admin_dash@vdart.com",
            password="Password123!",
            role="admin"
        )
        self.emp_day = CustomUser.objects.create_user(
            username="emp_day_1",
            email="emp_day@vdart.com",
            password="Password123!",
            role="employee",
            work_shift="day",
            monthly_tokens=100
        )
        self.emp_night = CustomUser.objects.create_user(
            username="emp_night_1",
            email="emp_night@vdart.com",
            password="Password123!",
            role="employee",
            work_shift="night",
            monthly_tokens=150
        )
        self.setting = CanteenSetting.get_settings()
        self.setting.token_allocation_day = 20
        self.setting.save()

    def test_dashboard_cycle_dates_adhere_to_20th_19th_policy(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/dashboard/stats/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        cycle_data = resp.data.get("cycleData")
        self.assertIsNotNone(cycle_data)
        
        from datetime import date
        c_start = date.fromisoformat(cycle_data["cycleStart"])
        c_end = date.fromisoformat(cycle_data["cycleEnd"])
        self.assertEqual(c_start.day, 20)
        self.assertEqual(c_end.day, 19)
        self.assertTrue(cycle_data["totalDays"] in (28, 29, 30, 31))
        self.assertTrue(cycle_data["daysElapsed"] >= 1)

    def test_cycle_revenue_data_endpoint_structure(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/admin/dashboard/cycle-revenue/?date=2026-10-22")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        c_info = resp.data.get("cycle_info")
        self.assertEqual(c_info["cycle_start"], "2026-10-20")
        self.assertEqual(c_info["cycle_end"], "2026-11-19")
        self.assertEqual(c_info["days_elapsed"], 3)
        self.assertEqual(c_info["total_days"], 31)
        
        shift_data = resp.data.get("shift_analytics")
        self.assertIn("day", shift_data)
        self.assertIn("night", shift_data)
        self.assertIn("balance", shift_data["day"])
        self.assertIn("allocated", shift_data["day"])

    def test_audit_export_view_filtering(self):
        self.client.force_authenticate(user=self.admin)
        resp = self.client.get("/api/reports/audit-export/?format=json&shift=day")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        allocs = resp.data.get("token_allocations", [])
        for a in allocs:
            self.assertEqual(a["shift"], "day")
        self.assertIn("available_cycles", resp.data)


# ---------------------------------------------------------
# 9. Multi-Office & Kitchen Workflows
# ---------------------------------------------------------
class OfficeAndKitchenTests(VDartCafeBaseTestCase):
    def setUp(self):
        super().setUp()
        self.trichy_office = Office.objects.create(
            name="Trichy Campus",
            code="TRZ",
            location="Trichy IT Park"
        )
        self.chennai_office = Office.objects.create(
            name="Chennai HQ",
            code="MAA",
            location="OMR Chennai"
        )
        self.trichy_kitchen = self.trichy_office.kitchen
        self.chennai_kitchen = self.chennai_office.kitchen

        # Create staff assigned to Trichy
        self.trichy_staff = CustomUser.objects.create_user(
            username="trichy_staff",
            user_id="STF_TRZ",
            email="staff.trz@vdart.com",
            password="StaffPassword123!",
            role="staff",
            assigned_kitchen=self.trichy_kitchen
        )

        # Create employee belonging to Trichy office
        self.trichy_emp = CustomUser.objects.create_user(
            username="trichy_emp",
            user_id="EMP_TRZ",
            email="emp.trz@vdart.com",
            password="EmpPassword123!",
            role="employee",
            work_shift="day",
            monthly_tokens=100,
            office=self.trichy_office
        )

        # Create employee belonging to Chennai office
        self.chennai_emp = CustomUser.objects.create_user(
            username="chennai_emp",
            user_id="EMP_MAA",
            email="emp.maa@vdart.com",
            password="EmpPassword123!",
            role="employee",
            work_shift="day",
            monthly_tokens=100,
            office=self.chennai_office
        )

        # Dishes in Trichy Kitchen
        self.dosa = MenuItem.objects.create(
            name="Trichy Special Dosa",
            kitchen=self.trichy_kitchen,
            price=10,
            is_available=True
        )

        # Dishes in Chennai Kitchen
        self.idli = MenuItem.objects.create(
            name="Chennai Sambar Idli",
            kitchen=self.chennai_kitchen,
            price=8,
            is_available=True
        )

    def test_office_creation_auto_creates_kitchen(self):
        """1-to-1 Mapping: creating an office must automatically create its paired kitchen."""
        blr_office = Office.objects.create(
            name="Bangalore Center",
            code="BLR",
            location="Electronic City"
        )
        self.assertTrue(hasattr(blr_office, 'kitchen'))
        self.assertEqual(blr_office.kitchen.name, "Bangalore Center Kitchen")
        self.assertTrue(blr_office.kitchen.is_active)

    def test_office_and_kitchen_admin_endpoints(self):
        """Admin can list, create, and fetch offices and kitchens via API."""
        self.client.force_authenticate(user=self.admin)
        
        # List offices
        resp = self.client.get("/api/admin/offices/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        office_names = [o['name'] for o in resp.data]
        self.assertIn("Trichy Campus", office_names)
        self.assertIn("Chennai HQ", office_names)

        # Create new office via API
        post_resp = self.client.post("/api/admin/offices/", {
            "name": "Hyderabad Hub",
            "code": "HYD",
            "location": "Hitec City"
        })
        self.assertEqual(post_resp.status_code, status.HTTP_201_CREATED)
        self.assertEqual(post_resp.data["code"], "HYD")
        self.assertIsNotNone(post_resp.data.get("kitchen_id"))

        # List kitchens
        kitchen_resp = self.client.get("/api/admin/kitchens/")
        self.assertEqual(kitchen_resp.status_code, status.HTTP_200_OK)
        kitchen_names = [k['name'] for k in kitchen_resp.data]
        self.assertTrue(any("Hyderabad" in name for name in kitchen_names))

    def test_menu_item_scoped_to_kitchen(self):
        """Employees only see menu items available in their office's kitchen."""
        # Trichy employee should see Dosa, but NOT Chennai's Idli
        self.client.force_authenticate(user=self.trichy_emp)
        resp = self.client.get("/api/employee/menu/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        item_names = [it['name'] for it in resp.data]
        self.assertIn("Trichy Special Dosa", item_names)
        self.assertNotIn("Chennai Sambar Idli", item_names)

        # Chennai employee should see Idli, but NOT Trichy's Dosa
        self.client.force_authenticate(user=self.chennai_emp)
        resp2 = self.client.get("/api/employee/menu/")
        self.assertEqual(resp2.status_code, status.HTTP_200_OK)
        chennai_item_names = [it['name'] for it in resp2.data]
        self.assertIn("Chennai Sambar Idli", chennai_item_names)
        self.assertNotIn("Trichy Special Dosa", chennai_item_names)

    def test_order_creation_routes_to_office_kitchen(self):
        """When an employee places an order, it is automatically routed to their office kitchen."""
        self.client.force_authenticate(user=self.trichy_emp)
        order_resp = self.client.post("/api/employee/order/", {
            "items": [{"menu_item_id": self.dosa.id, "quantity": 1}]
        }, format="json")
        self.assertEqual(order_resp.status_code, status.HTTP_201_CREATED)
        order_id = order_resp.data["id"]
        created_order = Order.objects.get(id=order_id)
        self.assertEqual(created_order.kitchen, self.trichy_kitchen)

        # Staff in Trichy sees this order
        self.client.force_authenticate(user=self.trichy_staff)
        staff_orders_resp = self.client.get("/api/staff/orders/")
        self.assertEqual(staff_orders_resp.status_code, status.HTTP_200_OK)
        staff_order_ids = [o['id'] for o in staff_orders_resp.data]
        self.assertIn(order_id, staff_order_ids)

    def test_staff_admin_and_staff_crew_isolation(self):
        """Staff Admin at Office A cannot view or manage Staff Admins or Staff Crew of Office B."""
        # Create Trichy Staff Admin
        trichy_mgr = CustomUser.objects.create_user(
            username="trichy_mgr",
            user_id="MGR_TRZ",
            email="mgr.trz@vdart.com",
            password="MgrPassword123!",
            role="staff_admin",
            assigned_kitchen=self.trichy_kitchen,
            office=self.trichy_office
        )
        # Create Chennai Staff Admin
        chennai_mgr = CustomUser.objects.create_user(
            username="chennai_mgr",
            user_id="MGR_MAA",
            email="mgr.maa@vdart.com",
            password="MgrPassword123!",
            role="staff_admin",
            assigned_kitchen=self.chennai_kitchen,
            office=self.chennai_office
        )
        # Create Chennai Staff Member
        chennai_staff = CustomUser.objects.create_user(
            username="chennai_staff",
            user_id="STF_MAA",
            email="staff.maa@vdart.com",
            password="StaffPassword123!",
            role="staff",
            assigned_kitchen=self.chennai_kitchen,
            office=self.chennai_office
        )

        # 1. Trichy staff_admin lists users
        self.client.force_authenticate(user=trichy_mgr)
        resp = self.client.get("/api/admin/users/")
        self.assertEqual(resp.status_code, status.HTTP_200_OK)
        usernames = [u["username"] for u in resp.data]
        
        # Must see Trichy staff
        self.assertIn("trichy_staff", usernames)
        # MUST NOT see Chennai staff
        self.assertNotIn("chennai_staff", usernames)
        # MUST NOT see other staff admins
        self.assertNotIn("chennai_mgr", usernames)
        self.assertNotIn("trichy_mgr", usernames)
        # MUST NOT see regular users (employees/guests)
        self.assertNotIn("trichy_emp", usernames)
        self.assertNotIn("chennai_emp", usernames)

        # 2. Trichy staff_admin cannot update or delete Chennai staff
        forbidden_update = self.client.patch(f"/api/admin/users/{chennai_staff.id}/", {"first_name": "Hacked"})
        self.assertEqual(forbidden_update.status_code, status.HTTP_404_NOT_FOUND)

        # 3. Super Admin can filter by category
        self.client.force_authenticate(user=self.admin)
        
        # Category: staff_admin
        sa_resp = self.client.get("/api/admin/users/?category=staff_admin")
        self.assertEqual(sa_resp.status_code, status.HTTP_200_OK)
        sa_names = [u["username"] for u in sa_resp.data]
        self.assertIn("trichy_mgr", sa_names)
        self.assertIn("chennai_mgr", sa_names)
        self.assertNotIn("trichy_staff", sa_names)
        self.assertNotIn("trichy_emp", sa_names)

        # Category: staff_crew
        sc_resp = self.client.get("/api/admin/users/?category=staff_crew")
        self.assertEqual(sc_resp.status_code, status.HTTP_200_OK)
        sc_names = [u["username"] for u in sc_resp.data]
        self.assertIn("trichy_staff", sc_names)
        self.assertIn("chennai_staff", sc_names)
        self.assertNotIn("trichy_mgr", sc_names)
        self.assertNotIn("trichy_emp", sc_names)

        # Category: users (employees and guests)
        users_resp = self.client.get("/api/admin/users/?category=users")
        self.assertEqual(users_resp.status_code, status.HTTP_200_OK)
        usr_names = [u["username"] for u in users_resp.data]
        self.assertIn("trichy_emp", usr_names)
        self.assertIn("chennai_emp", usr_names)
        self.assertNotIn("trichy_staff", usr_names)
        self.assertNotIn("trichy_mgr", usr_names)

        # Filter by Office
        trichy_only_resp = self.client.get(f"/api/admin/users/?office={self.trichy_office.id}")
        self.assertEqual(trichy_only_resp.status_code, status.HTTP_200_OK)
        trichy_only_names = [u["username"] for u in trichy_only_resp.data]
        self.assertIn("trichy_emp", trichy_only_names)
        self.assertNotIn("chennai_emp", trichy_only_names)




# ---------------------------------------------------------
# ---------------------------------------------------------
# 10. Audit Remediation Tests (C4, H1, H4, H9, M25)
# ---------------------------------------------------------
class AuditRemediationTests(VDartCafeBaseTestCase):
    def test_c4_staff_cannot_delete_orders(self):
        """C4: StaffOrderViewSet must reject DELETE on orders (405 Method Not Allowed)."""
        order = Order.objects.create(
            user=self.employee,
            total_tokens=15,
            status="pending"
        )
        self.client.force_authenticate(user=self.staff)
        resp = self.client.delete(f"/api/staff/orders/{order.id}/")
        self.assertEqual(resp.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)
        # Verify order still exists
        self.assertTrue(Order.objects.filter(id=order.id).exists())

    def test_h1_order_refund_is_atomic_and_idempotent(self):
        """H1: Order.refund_tokens locks row, credits tokens, and prevents double refunds."""
        initial_tokens = self.employee.monthly_tokens
        order = Order.objects.create(
            user=self.employee,
            total_tokens=25,
            status="pending"
        )
        # First refund succeeds
        success = order.refund_tokens()
        self.assertTrue(success)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, initial_tokens + 25)

        # Second refund is ignored (idempotent)
        second_success = order.refund_tokens()
        self.assertFalse(second_success)
        self.employee.refresh_from_db()
        self.assertEqual(self.employee.monthly_tokens, initial_tokens + 25)

    def test_h9_financial_records_immutable_via_api(self):
        """H9: Financial ledger must not allow DELETE or arbitrary POST mutations."""
        bill = CompanyBill.objects.create(
            bill_number="BILL-TEST-001",
            billing_period_start=timezone.now().date(),
            billing_period_end=timezone.now().date(),
            total_tokens_billed=100,
            status="generated"
        )
        tx = FinancialTransaction.objects.create(
            user=self.employee,
            transaction_type="deduction",
            amount_tokens=50
        )
        self.client.force_authenticate(user=self.admin)

        # Company bills reject DELETE
        bill_del = self.client.delete(f"/api/finance/company-bills/{bill.id}/")
        self.assertEqual(bill_del.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)

        # Financial transactions reject DELETE
        tx_del = self.client.delete(f"/api/finance/transactions/{tx.id}/")
        self.assertEqual(tx_del.status_code, status.HTTP_405_METHOD_NOT_ALLOWED)

    def test_h4_session_cookie_remember_me_respected(self):
        """H4: Login with remember_me=False must not set max_age on session cookie."""
        from django.conf import settings
        session_cookie = settings.SESSION_COOKIE_NAME
        resp_no_remember = self.client.post("/api/login/", {
            "username": "emp_test",
            "password": "EmpPassword123!",
            "remember_me": False
        }, format="json")
        self.assertEqual(resp_no_remember.status_code, status.HTTP_200_OK)
        cookie = resp_no_remember.cookies.get(session_cookie)
        self.assertIsNotNone(cookie)
        self.assertEqual(cookie["max-age"], "")

        # When remember_me=True, max-age is set to 2 weeks
        resp_remember = self.client.post("/api/login/", {
            "username": "emp_test",
            "password": "EmpPassword123!",
            "remember_me": True
        }, format="json")
        self.assertEqual(resp_remember.status_code, status.HTTP_200_OK)
        cookie_rem = resp_remember.cookies.get(session_cookie)
        self.assertEqual(cookie_rem["max-age"], 1209600)

    def test_m25_staff_admin_cannot_inspect_other_kitchen_or_superadmin(self):
        """M25: Staff Admin cannot access token transactions of Admins or other offices."""
        other_office = Office.objects.create(name="Other HQ", code="OTH")
        other_emp = CustomUser.objects.create_user(
            username="other_emp",
            user_id="EMP_OTH",
            email="other.emp@vdart.com",
            password="EmpPassword123!",
            role="employee",
            office=other_office
        )
        my_office = Office.objects.create(name="My Office", code="MYO")
        my_staff_admin = CustomUser.objects.create_user(
            username="my_mgr",
            user_id="MGR_MYO",
            email="mgr.myo@vdart.com",
            password="MgrPassword123!",
            role="staff_admin",
            office=my_office,
            assigned_kitchen=my_office.kitchen
        )
        self.client.force_authenticate(user=my_staff_admin)

        # Attempt to inspect Super Admin -> 403 Forbidden
        admin_tx = self.client.get(f"/api/admin/tokens/user-transactions/{self.admin.id}/")
        self.assertEqual(admin_tx.status_code, status.HTTP_403_FORBIDDEN)

        # Attempt to inspect employee of another office -> 403 Forbidden
        other_tx = self.client.get(f"/api/admin/tokens/user-transactions/{other_emp.id}/")
        self.assertEqual(other_tx.status_code, status.HTTP_403_FORBIDDEN)

    def test_h3_auth_rate_limiting_configured(self):
        """H3: Public auth endpoints must use AuthAttemptThrottle with 10/min scope."""
        from api.views import login_view, reset_password_view, submit_password_reset_request, AuthAttemptThrottle
        for view_func in (login_view, reset_password_view, submit_password_reset_request):
            throttle_classes = getattr(view_func, 'cls', None).throttle_classes
            self.assertIn(AuthAttemptThrottle, throttle_classes)
        self.assertEqual(AuthAttemptThrottle.scope, 'auth_attempt')

    def test_h8_cors_preflight_and_csrf_clean(self):
        """H8: CSRF and Auth preflight work cleanly without conflicting manual headers."""
        csrf_resp = self.client.get("/api/csrf/")
        self.assertEqual(csrf_resp.status_code, status.HTTP_200_OK)
        self.assertIn("csrftoken", csrf_resp.data)

        options_resp = self.client.options("/api/login/")
        self.assertEqual(options_resp.status_code, status.HTTP_200_OK)
